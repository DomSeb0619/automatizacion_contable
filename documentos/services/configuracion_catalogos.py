"""Previsualizacion y aplicacion transaccional de catalogos desde la interfaz."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from django.db import transaction

from documentos.models import ActualizacionCatalogo, MapeoProductoIVA, MapeoRubroCuenta, Proyecto, normalize_code
from documentos.services.categorias_predeterminadas import categoria_predeterminada


PRODUCT_HEADERS = {"CODPROD01", "DESPROD01", "PORCIVA01"}
RUBRO_HEADERS = {"CODIGO DE PROYECTO", "CODIGO DEL RUBRO PADRE", "CODIGO DE RUBRO", "DESCRIPCION DEL RUBRO", "TIPO", "CUENTAS"}


class CatalogoInvalido(ValueError):
    """El archivo no tiene la estructura necesaria para ser importado."""


class ConfirmacionInvalida(ValueError):
    """La previsualizacion no puede persistirse de forma segura."""


@dataclass(frozen=True)
class VistaCatalogo:
    tipo: str
    proyecto_id: int
    filas: list[dict]
    creados: int
    actualizados: int
    omitidos: int
    rechazados: int
    errores: list[str]
    nombre_archivo: str = ""

    @property
    def puede_confirmar(self) -> bool:
        return not self.errores

    @property
    def requiere_aceptacion(self) -> bool:
        return self.rechazados > 0

    def as_session(self) -> dict:
        return {
            "tipo": self.tipo,
            "proyecto_id": self.proyecto_id,
            "filas": self.filas,
            "creados": self.creados,
            "actualizados": self.actualizados,
            "omitidos": self.omitidos,
            "rechazados": self.rechazados,
            "errores": self.errores,
            "nombre_archivo": self.nombre_archivo,
        }


def previsualizar_catalogo(
    proyecto: Proyecto, tipo: str, source: str | Path, *, nombre_archivo: str = ""
) -> VistaCatalogo:
    """Lee un archivo temporal y clasifica cambios sin escribir en SQLite."""
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as error:
        raise CatalogoInvalido("No se puede leer el archivo seleccionado.") from error
    try:
        if tipo == "productos":
            sheet = workbook.active
            required = PRODUCT_HEADERS
        elif tipo == "rubros":
            if "Hoja2" not in workbook.sheetnames:
                raise CatalogoInvalido("El catalogo de rubros debe contener la hoja Hoja2.")
            sheet = workbook["Hoja2"]
            required = RUBRO_HEADERS
        else:
            raise CatalogoInvalido("Tipo de catalogo no reconocido.")

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise CatalogoInvalido("El archivo no contiene encabezados.")
        headers = {str(value).strip().upper(): index for index, value in enumerate(rows[0]) if value is not None}
        if not required.issubset(headers):
            raise CatalogoInvalido("Los encabezados del archivo no corresponden al catalogo seleccionado.")
        if tipo == "productos":
            return preview_productos(proyecto, rows[1:], headers, nombre_archivo=nombre_archivo)
        validate_rubro_project_scope(proyecto, rows[1:], headers)
        return preview_rubros(proyecto, rows[1:], headers, nombre_archivo=nombre_archivo)
    finally:
        workbook.close()


def preview_productos(proyecto: Proyecto, rows: list[tuple], headers: dict[str, int], *, nombre_archivo: str = "") -> VistaCatalogo:
    existing = {item.codigo_producto: item for item in MapeoProductoIVA.objects.filter(proyecto=proyecto)}
    seen: set[str] = set()
    preview: list[dict] = []
    counts = Counter()
    for row_number, row in enumerate(rows, start=2):
        code = normalize_code(value_at(row, headers["CODPROD01"]))
        description = str(value_at(row, headers["DESPROD01"]) or "").strip()
        raw_rate = value_at(row, headers["PORCIVA01"])
        rate = decimal_rate(raw_rate)
        reason = validate_product_row(code, rate, seen)
        if reason:
            preview.append({"fila": row_number, "codigo": code, "estado": "RECHAZADO", "motivo": reason, "valor_problematico": raw_rate if raw_rate is not None else code})
            counts.rejected += 1
            continue
        seen.add(code)
        current = existing.get(code)
        values = {
            "descripcion": description,
            "porcentaje_iva": f"{rate:.4f}",
            "activo": True,
            # Solo proponemos el valor por prefijo cuando la categoria aun esta vacia.
            "categoria": current.categoria if current and current.categoria else categoria_predeterminada(code),
        }
        state = "CREAR" if current is None else "ACTUALIZAR" if product_changed(current, values) else "OMITIR"
        counts.add(state)
        preview.append({"fila": row_number, "codigo": code, "estado": state, "descripcion": description, "porcentaje_iva": values["porcentaje_iva"], "categoria": values["categoria"]})
    return make_preview("productos", proyecto, preview, counts, nombre_archivo=nombre_archivo)


def preview_rubros(proyecto: Proyecto, rows: list[tuple], headers: dict[str, int], *, nombre_archivo: str = "") -> VistaCatalogo:
    existing = {item.codigo_rubro: item for item in MapeoRubroCuenta.objects.filter(proyecto=proyecto)}
    seen: set[str] = set()
    preview: list[dict] = []
    counts = Counter()
    for row_number, row in enumerate(rows, start=2):
        code = normalize_code(value_at(row, headers["CODIGO DE RUBRO"]))
        description = str(value_at(row, headers["DESCRIPCION DEL RUBRO"]) or "").strip()
        tipo = str(value_at(row, headers["TIPO"]) or "").strip().upper()
        account = str(value_at(row, headers["CUENTAS"]) or "").strip().upper()
        reason = validate_rubro_row(code, tipo, account, seen)
        if reason:
            if not code or code in seen:
                problem_value = code
            elif tipo not in {"P", "V"}:
                problem_value = tipo
            else:
                problem_value = account
            preview.append({"fila": row_number, "codigo": code, "estado": "RECHAZADO", "motivo": reason, "valor_problematico": problem_value or code})
            counts.rejected += 1
            continue
        seen.add(code)
        values = {"descripcion": description, "tipo": tipo, "cuenta_contable": account, "activo": True}
        current = existing.get(code)
        state = "CREAR" if current is None else "ACTUALIZAR" if rubro_changed(current, values) else "OMITIR"
        counts.add(state)
        preview.append({"fila": row_number, "codigo": code, "estado": state, **values})
    return make_preview("rubros", proyecto, preview, counts, nombre_archivo=nombre_archivo)


def aplicar_previsualizacion(payload: dict, *, acepta_rechazados: bool = False) -> VistaCatalogo:
    """Aplica un borrador de sesion. Las filas rechazadas nunca se persisten."""
    if payload.get("errores"):
        raise ConfirmacionInvalida("La previsualizacion contiene errores estructurales y no puede confirmarse.")
    if payload.get("rechazados") and not acepta_rechazados:
        raise ConfirmacionInvalida("Confirma que las filas rechazadas seran excluidas antes de continuar.")
    try:
        proyecto = Proyecto.objects.select_related("empresa").get(pk=payload["proyecto_id"], activo=True, empresa__activa=True)
    except (KeyError, Proyecto.DoesNotExist) as error:
        raise ConfirmacionInvalida("El proyecto seleccionado ya no esta disponible.") from error
    tipo = payload.get("tipo")
    rows = payload.get("filas", [])
    counts = Counter()
    counts.rejected = payload.get("rechazados", 0)
    with transaction.atomic():
        for row in rows:
            if row.get("estado") == "OMITIR":
                counts.skipped += 1
                continue
            if row.get("estado") == "RECHAZADO":
                continue
            if tipo == "productos":
                apply_producto(proyecto, row, counts)
            elif tipo == "rubros":
                apply_rubro(proyecto, row, counts)
            else:
                raise ConfirmacionInvalida("Tipo de catalogo no reconocido.")
        ActualizacionCatalogo.objects.update_or_create(
            proyecto=proyecto,
            tipo=ActualizacionCatalogo.Tipo.PRODUCTOS if tipo == "productos" else ActualizacionCatalogo.Tipo.RUBROS,
            defaults={
                "nombre_archivo": Path(str(payload.get("nombre_archivo", ""))).name,
                "creados": counts.created,
                "actualizados": counts.updated,
                "omitidos": counts.skipped,
                "rechazados": counts.rejected,
            },
        )
    return make_preview(tipo, proyecto, rows, counts, nombre_archivo=payload.get("nombre_archivo", ""))


def apply_producto(proyecto: Proyecto, row: dict, counts: "Counter") -> None:
    defaults = {
        "descripcion": row["descripcion"],
        "porcentaje_iva": Decimal(row["porcentaje_iva"]),
        "activo": True,
        "categoria": row.get("categoria", categoria_predeterminada(row["codigo"])),
    }
    current = MapeoProductoIVA.objects.filter(proyecto=proyecto, codigo_producto=row["codigo"]).first()
    if current is None:
        MapeoProductoIVA.objects.create(proyecto=proyecto, codigo_producto=row["codigo"], **defaults)
        counts.created += 1
    elif product_changed(current, {**defaults, "porcentaje_iva": f"{defaults['porcentaje_iva']:.4f}"}):
        for field, value in defaults.items():
            setattr(current, field, value)
        current.save()
        counts.updated += 1
    else:
        counts.skipped += 1


def apply_rubro(proyecto: Proyecto, row: dict, counts: "Counter") -> None:
    defaults = {"descripcion": row["descripcion"], "tipo": row["tipo"], "cuenta_contable": row["cuenta_contable"], "activo": True}
    current = MapeoRubroCuenta.objects.filter(proyecto=proyecto, codigo_rubro=row["codigo"]).first()
    if current is None:
        MapeoRubroCuenta.objects.create(proyecto=proyecto, codigo_rubro=row["codigo"], **defaults)
        counts.created += 1
    elif rubro_changed(current, defaults):
        for field, value in defaults.items():
            setattr(current, field, value)
        current.save()
        counts.updated += 1
    else:
        counts.skipped += 1


def validate_product_row(code: str, rate: Decimal | None, seen: set[str]) -> str | None:
    if not code:
        return "Codigo de producto vacio."
    if code in seen:
        return "Codigo de producto duplicado en el archivo."
    if rate is None or not Decimal("0") <= rate <= Decimal("100"):
        return "IVA invalido o fuera del rango 0-100."
    return None


def validate_rubro_row(code: str, tipo: str, account: str, seen: set[str]) -> str | None:
    if not code:
        return "Codigo de rubro vacio."
    if code in seen:
        return "Codigo de rubro duplicado en el archivo."
    if tipo not in {"P", "V"}:
        return "Tipo de rubro invalido."
    if not account:
        return "Cuenta contable vacia."
    return None


def validate_rubro_project_scope(proyecto: Proyecto, rows: list[tuple], headers: dict[str, int]) -> None:
    projects = {normalize_code(value_at(row, headers["CODIGO DE PROYECTO"])) for row in rows if value_at(row, headers["CODIGO DE PROYECTO"])}
    if projects and projects != {normalize_code(proyecto.codigo)}:
        raise CatalogoInvalido("El archivo de rubros no corresponde al proyecto seleccionado.")


def decimal_rate(value) -> Decimal | None:
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, AttributeError):
        return None


def value_at(row: tuple, index: int):
    return row[index] if index < len(row) else None


def product_changed(current: MapeoProductoIVA, values: dict) -> bool:
    return (
        current.descripcion != values["descripcion"]
        or current.porcentaje_iva != Decimal(values["porcentaje_iva"])
        or current.activo != values["activo"]
        or current.categoria != values.get("categoria", current.categoria)
    )


def rubro_changed(current: MapeoRubroCuenta, values: dict) -> bool:
    return any(getattr(current, field) != value for field, value in values.items())


def make_preview(
    tipo: str, proyecto: Proyecto, rows: list[dict], counts: "Counter", *, nombre_archivo: str = ""
) -> VistaCatalogo:
    return VistaCatalogo(tipo, proyecto.pk, rows, counts.created, counts.updated, counts.skipped, counts.rejected, [], nombre_archivo)


class Counter:
    def __init__(self) -> None:
        self.created = 0
        self.updated = 0
        self.skipped = 0
        self.rejected = 0

    def add(self, state: str) -> None:
        if state == "CREAR":
            self.created += 1
        elif state == "ACTUALIZAR":
            self.updated += 1
        else:
            self.skipped += 1
