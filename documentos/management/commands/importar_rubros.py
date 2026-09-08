from __future__ import annotations

from openpyxl import load_workbook

from documentos.management.commands._catalogo_base import CatalogoCommand, ImportStats
from documentos.models import MapeoRubroCuenta, normalize_code


EXPECTED_HEADERS = {
    "CODIGO DE PROYECTO",
    "CODIGO DEL RUBRO PADRE",
    "CODIGO DE RUBRO",
    "DESCRIPCION DEL RUBRO",
    "TIPO",
    "CUENTAS",
}


class Command(CatalogoCommand):
    help = "Importa Hoja2 del catalogo de rubros y cuentas."

    def add_arguments(self, parser) -> None:
        parser.add_argument("path")
        self.add_scope_arguments(parser)

    def handle(self, *args, **options):
        source = self.validate_source(options["path"])

        def operation():
            empresa, proyecto = self.get_scope(options)
            workbook = load_workbook(source, read_only=True, data_only=True)
            if "Hoja2" not in workbook.sheetnames:
                self.stderr.write("ERROR: falta la hoja Hoja2.")
                return ImportStats(errors=1)
            sheet = workbook["Hoja2"]
            headers = {str(value).strip().upper() for value in next(sheet.iter_rows(values_only=True)) if value is not None}
            if not EXPECTED_HEADERS.issubset(headers):
                self.stderr.write("ERROR: encabezados invalidos en Hoja2.")
                return ImportStats(errors=1)

            stats = ImportStats()
            seen: set[str] = set()
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                code = normalize_code(row[2] or "")
                description = str(row[3] or "").strip()
                tipo = str(row[4] or "").strip().upper()
                account = str(row[5] or "").strip().upper()
                if not code:
                    stats.errors += 1
                    self.stderr.write(f"Fila {row_number}: codigo de rubro vacio.")
                    continue
                if code in seen:
                    stats.duplicates += 1
                    self.stderr.write(f"Fila {row_number}: rubro duplicado {code}.")
                    continue
                seen.add(code)
                if tipo not in {"P", "V"}:
                    stats.errors += 1
                    self.stderr.write(f"Fila {row_number}: tipo invalido {tipo!r}.")
                    continue
                if not account:
                    stats.errors += 1
                    self.stderr.write(f"Fila {row_number}: cuenta vacia para rubro {code}.")
                    continue
                existing = MapeoRubroCuenta.objects.filter(proyecto=proyecto, codigo_rubro=code).first()
                values = {"descripcion": description, "tipo": tipo, "cuenta_contable": account, "activo": True}
                if existing is None:
                    MapeoRubroCuenta.objects.create(proyecto=proyecto, codigo_rubro=code, **values)
                    stats.created += 1
                elif any(getattr(existing, field) != value for field, value in values.items()):
                    for field, value in values.items():
                        setattr(existing, field, value)
                    existing.save()
                    stats.updated += 1
                else:
                    stats.skipped += 1
            return stats

        stats = self.run_atomic(options["dry_run"], operation)
        self.write_stats(stats)
