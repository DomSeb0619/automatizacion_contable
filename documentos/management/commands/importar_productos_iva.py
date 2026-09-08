from __future__ import annotations

from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from documentos.management.commands._catalogo_base import CatalogoCommand, ImportStats
from documentos.models import MapeoProductoIVA, normalize_code
from documentos.services.categorias_predeterminadas import categoria_predeterminada


EXPECTED_HEADERS = {"CODPROD01", "DESPROD01", "PORCIVA01"}


class Command(CatalogoCommand):
    help = "Importa el maestro de productos e IVA."

    def add_arguments(self, parser) -> None:
        parser.add_argument("path")
        self.add_scope_arguments(parser)

    def handle(self, *args, **options):
        source = self.validate_source(options["path"])

        def operation():
            _, proyecto = self.get_scope(options)
            workbook = load_workbook(source, read_only=True, data_only=True)
            sheet = workbook.active
            headers = {str(value).strip().upper() for value in next(sheet.iter_rows(values_only=True)) if value is not None}
            if not EXPECTED_HEADERS.issubset(headers):
                self.stderr.write("ERROR: encabezados invalidos en maestro de productos.")
                return ImportStats(errors=1)

            stats = ImportStats()
            seen: set[str] = set()
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                code = normalize_code(row[0] or "")
                description = str(row[1] or "").strip()
                try:
                    rate = Decimal(str(row[2]).strip())
                except (InvalidOperation, AttributeError):
                    stats.errors += 1
                    self.stderr.write(f"Fila {row_number}: IVA invalido para {code or 'sin codigo'}.")
                    continue
                if not code:
                    stats.errors += 1
                    self.stderr.write(f"Fila {row_number}: codigo de producto vacio.")
                    continue
                if code in seen:
                    stats.duplicates += 1
                    self.stderr.write(f"Fila {row_number}: producto duplicado {code}.")
                    continue
                seen.add(code)
                if not Decimal("0") <= rate <= Decimal("100"):
                    stats.errors += 1
                    self.stderr.write(f"Fila {row_number}: IVA fuera de rango 0-100 para {code}: {rate}.")
                    continue
                existing = MapeoProductoIVA.objects.filter(proyecto=proyecto, codigo_producto=code).first()
                values = {"descripcion": description, "porcentaje_iva": rate, "activo": True}
                if existing is None:
                    MapeoProductoIVA.objects.create(proyecto=proyecto, codigo_producto=code, categoria=categoria_predeterminada(code), **values)
                    stats.created += 1
                elif any(getattr(existing, field) != value for field, value in values.items()):
                    for field, value in values.items():
                        setattr(existing, field, value)
                    existing.save()
                    stats.updated += 1
                else:
                    stats.skipped += 1
            return stats

        stats = self.run_atomic(options["dry_run"], operation)
        self.write_stats(stats)
